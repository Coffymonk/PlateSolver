# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright © 2026 Miklos Elmberg
"""Make a quiz file (.dat) from a spreadsheet the user filled in: one row per question.

Columns (the first row holds the headings; their order doesn't matter):
    Topic | Level | Question | Right answer | Wrong answer 1 … Wrong answer 4 | Explanation | Picture

Excel (.xlsx) and CSV files are read; CSV may use commas, semicolons (Swedish and other European Excel)
or tabs. Every problem is reported in plain words with its row number, so it's easy to fix.
"""
from __future__ import annotations

import csv
import datetime
import hashlib
import io
import math
import re
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np

from platesolver.core.quiz import Question, QuizBank, Topic, my_quizzes_dir, question_id

COLUMNS = ["Topic", "Level", "Question", "Right answer", "Wrong answer 1", "Wrong answer 2", "Wrong answer 3",
           "Wrong answer 4", "Explanation", "Picture"]
EXAMPLES = [
    ["The Moon", 1, "How long does the Moon take to orbit the Earth (relative to the stars)?", "About 27 days",
     "About 7 days", "About 365 days", "About 2 days", "", "The sidereal month is about 27.3 days; the cycle of "
     "phases (29.5 days) is longer because the Earth moves around the Sun meanwhile.", ""],
    ["The Moon", 2, "Why do we always see the same side of the Moon?", "Its rotation is locked to its orbit",
     "It doesn't rotate at all", "The far side is always dark", "", "", "Tidal forces slowed the Moon's spin "
     "until one rotation took exactly one orbit (tidal locking).", ""],
    ["My telescope", 1, "What does the focal ratio f/5 mean?", "Focal length is 5 times the aperture",
     "Aperture is 5 times the focal length", "The telescope magnifies 5 times", "The eyepiece is 5 mm", "",
     "f-number = focal length ÷ aperture, e.g. 500 mm ÷ 100 mm = f/5.", ""],
]
HELP_LINES = [
    "How to make your own PlateSolver quiz",
    "",
    "1. Fill in the sheet 'Questions': one row per question. Replace or delete the example rows.",
    "2. Topic: any name, e.g. 'The Moon'. Questions with the same topic are grouped (choose a topic in the quiz).",
    "3. Level: 1 = beginner, 2 = intermediate, 3 = advanced. Empty means 1.",
    "4. Question, Right answer and at least two Wrong answers are needed (up to four). The answers are shuffled.",
    "5. Explanation (optional): shown after answering.",
    "6. Picture (optional): the file name of an image next to this spreadsheet, used for that topic.",
    "   Without one, PlateSolver draws a starry picture for each topic.",
    "7. Save the file, then in PlateSolver: Help › Astronomy quiz › Create a quiz… and choose this file.",
    "",
    "Tip: avoid answers like 'All of the above' – the answers are shown in a random order.",
]

_ALIASES = {
    "topic": ("topic", "category", "subject"),
    "level": ("level", "difficulty"),
    "question": ("question", "q"),
    "right": ("right answer", "correct answer", "right", "correct", "answer"),
    "explain": ("explanation", "explain", "why"),
    "picture": ("picture", "image", "photo"),
}


@dataclass
class QuizDraft:
    """What was read from the spreadsheet: the quiz (if usable), and any problems."""
    bank: QuizBank | None = None
    errors: list[str] = field(default_factory=list)       # must be fixed
    warnings: list[str] = field(default_factory=list)     # worth a look
    rows: int = 0

    @property
    def ok(self) -> bool:
        return self.bank is not None and not self.errors and bool(self.bank.questions)

    def summary(self) -> str:
        if self.bank is None or not self.bank.questions:
            return "No questions found."
        topics = len(self.bank.topics)
        return f"{len(self.bank.questions)} questions in {topics} topic{'s' if topics != 1 else ''}"


# --------------------------------------------------------------------------- the template
def write_template(path: Path) -> Path:
    """A ready-to-fill spreadsheet with headings, example rows and instructions."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(COLUMNS)
            w.writerows(EXAMPLES)
        return path
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(COLUMNS)
    for row in EXAMPLES:
        ws.append(row)
    widths = [18, 7, 50, 30, 26, 26, 26, 26, 50, 16]
    head = PatternFill("solid", fgColor="25405F")
    for i, wdt in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
        ws.column_dimensions[cell.column_letter].width = wdt
    for row in ws.iter_rows(min_row=2, max_row=200):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True,
                        error="Level must be 1 (beginner), 2 (intermediate) or 3 (advanced).")
    ws.add_data_validation(dv)
    dv.add("B2:B2000")
    info = wb.create_sheet("How to")
    for line in HELP_LINES:
        info.append([line])
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 110
    wb.save(path)
    return path


# --------------------------------------------------------------------------- reading
def _read_table(path: Path) -> list[list[str]]:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = [wb["Questions"]] if "Questions" in wb.sheetnames else []
        sheets += [ws for ws in wb.worksheets if ws not in sheets]
        for ws in sheets:
            rows = [["" if v is None else str(v).strip() for v in row] for row in ws.iter_rows(values_only=True)]
            if _header_index(rows) is not None:
                return rows
        return []
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    first = text.splitlines()[0] if text else ""
    delim = max((";", ",", "\t"), key=first.count)
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delim)]


def _header_index(rows: list[list[str]]) -> int | None:
    for i, row in enumerate(rows[:20]):
        names = [c.strip().lower() for c in row]
        if "question" in names or "q" in names:
            return i
    return None


def _columns(header: list[str]) -> dict[str, int | list[int]]:
    names = [c.strip().lower() for c in header]
    cols: dict = {"wrong": []}
    for key, aliases in _ALIASES.items():
        for a in aliases:
            if a in names:
                cols[key] = names.index(a)
                break
    for i, n in enumerate(names):
        if n.startswith("wrong") or n.startswith("incorrect"):
            cols["wrong"].append(i)
    return cols


def _cell(row: list[str], i) -> str:
    return row[i].strip() if isinstance(i, int) and i < len(row) and row[i] is not None else ""


def read_spreadsheet(path: Path, title: str = "") -> QuizDraft:
    """Read and check a filled-in spreadsheet. Nothing is written."""
    path = Path(path)
    draft = QuizDraft()
    try:
        rows = _read_table(path)
    except Exception as exc:
        draft.errors.append(f"The file could not be read: {exc}")
        return draft
    h = _header_index(rows)
    if h is None:
        draft.errors.append("No heading row with a 'Question' column was found. Start from the template "
                            "(Get the template) and keep its first row.")
        return draft
    cols = _columns(rows[h])
    for need, label in (("question", "Question"), ("right", "Right answer")):
        if need not in cols:
            draft.errors.append(f"The column '{label}' is missing from the heading row.")
    if len(cols["wrong"]) < 2:
        draft.errors.append("At least two 'Wrong answer' columns are needed (Wrong answer 1, Wrong answer 2, …).")
    if draft.errors:
        return draft

    topics: dict[str, Topic] = {}
    pictures: dict[str, Path] = {}
    questions: list[Question] = []
    seen: set[str] = set()
    for n, row in enumerate(rows[h + 1:], start=h + 2):           # spreadsheet row numbers
        if not any(c.strip() for c in row if c):
            continue
        draft.rows += 1
        q_text = _cell(row, cols["question"])
        right = _cell(row, cols["right"])
        wrong = [w for w in (_cell(row, i) for i in cols["wrong"]) if w]
        topic_name = _cell(row, cols.get("topic")) or "General"
        level_text = _cell(row, cols.get("level"))
        problems = []
        if not q_text:
            problems.append("the question is empty")
        if not right:
            problems.append("the right answer is empty")
        if len(wrong) < 2:
            problems.append(f"needs at least two wrong answers (has {len(wrong)})")
        if len(wrong) > 4:
            problems.append("has more than four wrong answers")
        options = [right] + wrong
        if right and len({o.lower() for o in options}) < len(options):
            problems.append("two answers are the same")
        level = 1
        if level_text:
            try:
                level = int(float(level_text.replace(",", ".")))
            except ValueError:
                level = 0
            if level not in (1, 2, 3):
                problems.append(f"level must be 1, 2 or 3 (it says '{level_text}')")
        if problems:
            draft.errors.append(f"Row {n}: " + "; ".join(problems) + ".")
            continue
        if q_text.lower() in seen:
            draft.warnings.append(f"Row {n}: the same question appears earlier; only the first is used.")
            continue
        seen.add(q_text.lower())
        if any(re.search(r"\b(all|none|both) of (the )?(above|these)\b", o, re.I) for o in options):
            draft.warnings.append(f"Row {n}: answers like 'all of the above' don't work well, because the "
                                  "answers are shown in a random order.")
        tid = _topic_id(topic_name)
        if tid not in topics:
            topics[tid] = Topic(tid, topic_name)
        pic = _cell(row, cols.get("picture"))
        if pic and tid not in pictures:
            p = Path(pic) if Path(pic).is_absolute() else path.parent / pic
            if p.is_file():
                pictures[tid] = p
            else:
                draft.warnings.append(f"Row {n}: the picture '{pic}' was not found next to the spreadsheet; "
                                      "a drawn picture is used instead.")
        questions.append(Question(question_id(q_text), tid, q_text, options, 0,
                                  _cell(row, cols.get("explain")), level))
    if not questions and not draft.errors:
        draft.errors.append("No questions were found below the heading row.")
    for tid, t in topics.items():
        t.image = _picture_bytes(pictures.get(tid), t.name, draft.warnings)
    draft.bank = QuizBank(title or path.stem, datetime.date.today().isoformat(), list(topics.values()),
                          questions, str(path))
    return draft


def _topic_id(name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")[:30] or "topic"
    return f"{slug}_{hashlib.sha1(name.encode('utf-8')).hexdigest()[:4]}"


# --------------------------------------------------------------------------- pictures
def _picture_bytes(path: Path | None, topic: str, warnings: list[str]) -> bytes:
    from PIL import Image
    img = None
    if path is not None:
        try:
            with Image.open(path) as im:
                img = _fit(im.convert("RGB"))
        except Exception as exc:
            warnings.append(f"The picture {path.name} could not be read ({exc}); a drawn picture is used instead.")
    if img is None:
        img = auto_picture(topic)
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=True)
    return buf.getvalue()


def _fit(im, w: int = 240, h: int = 480):
    """Crop to 1:2 (twice as high as wide) around the centre, then scale."""
    from PIL import Image
    sw, sh = im.size
    target = w / h
    if sw / sh > target:
        nw = int(sh * target)
        im = im.crop(((sw - nw) // 2, 0, (sw - nw) // 2 + nw, sh))
    else:
        nh = int(sw / target)
        im = im.crop((0, (sh - nh) // 2, sw, (sh - nh) // 2 + nh))
    return im.resize((w, h), Image.LANCZOS)


def auto_picture(topic: str, w: int = 240, h: int = 480):
    """A starry picture with a soft coloured glow; the colour and layout come from the topic's name."""
    from PIL import Image, ImageDraw, ImageFilter
    seed = int(hashlib.sha1(topic.encode("utf-8")).hexdigest()[:8], 16)
    rng = np.random.default_rng(seed)
    t = np.linspace(0, 1, h)[:, None, None]
    img = np.array([6, 9, 20.0]) * (1 - t) + np.array([20, 28, 54.0]) * t
    img = np.repeat(img, w, axis=1)
    hue = rng.uniform(0, 1)
    col = np.array([0.5 + 0.5 * math.cos(2 * math.pi * (hue + k / 3)) for k in range(3)]) * 200
    yy, xx = np.mgrid[0:h, 0:w]
    for _ in range(3):
        cx, cy, r = rng.uniform(40, w - 40), rng.uniform(80, h - 80), rng.uniform(35, 90)
        img += np.exp(-((xx - cx) ** 2 + (yy - cy) ** 2) / (2 * r * r))[..., None] * col * rng.uniform(0.3, 0.7)
    im = Image.fromarray(np.clip(img + rng.normal(0, 1.2, img.shape), 0, 255).astype(np.uint8))
    d = ImageDraw.Draw(im)
    for _ in range(260):
        x, y = rng.uniform(0, w), rng.uniform(0, h)
        b = rng.power(3)
        r = 0.5 + rng.power(8) * 1.6
        d.ellipse([x - r, y - r, x + r, y + r], fill=tuple(int(255 * b * c) for c in (1.0, 0.95, 0.9)))
    return im.filter(ImageFilter.GaussianBlur(0.4))


# --------------------------------------------------------------------------- saving
def safe_file_name(title: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "", title).strip().rstrip(".") or "My quiz"
    return name[:80] + ".dat"


def save_quiz(draft: QuizDraft, title: str, folder: Path | None = None) -> Path:
    """Write the quiz file (in My quizzes by default). Returns its path."""
    if not draft.ok:
        raise ValueError("the spreadsheet still has problems")
    draft.bank.title = title.strip() or draft.bank.title
    folder = folder or my_quizzes_dir()
    path = Path(folder) / safe_file_name(draft.bank.title)
    draft.bank.save(path)
    return path
